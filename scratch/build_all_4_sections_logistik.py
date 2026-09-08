import openpyxl, json

wb = openpyxl.load_workbook('2026-LAP LOGISTIK.xlsx', data_only=True)
db_all = {}

for name in ['JAN', 'FEB', 'MAR', 'APRIL', 'MEI', 'JUNI']:
    ws = wb[name]
    
    # Sec II (Penjualan)
    sec2 = []
    for r in range(17, 23):
        val_name = ws.cell(r, 3).value
        if val_name and val_name != 'JUMLAH TOTAL':
            sec2.append({
                'no': ws.cell(r, 2).value,
                'nama': str(val_name).strip(),
                'tunai_kg': ws.cell(r, 4).value or 0,
                'tunai_harga': ws.cell(r, 5).value or 0,
                'tunai_rp': ws.cell(r, 6).value or 0,
                'pot_kg': ws.cell(r, 7).value or 0,
                'pot_harga': ws.cell(r, 8).value or 0,
                'pot_rp': ws.cell(r, 9).value or 0,
                'piu_kg': ws.cell(r, 10).value or 0,
                'piu_harga': ws.cell(r, 11).value or 0,
                'piu_rp': ws.cell(r, 12).value or 0,
                'bunt_kg': ws.cell(r, 13).value or 0,
                'bunt_harga': ws.cell(r, 14).value or 0,
                'bunt_rp': ws.cell(r, 15).value or 0,
                'total_kg': ws.cell(r, 16).value or 0,
                'total_rp': ws.cell(r, 17).value or 0,
            })
            
    # Sec III (Pembelian)
    sec3 = []
    for r in range(28, 33):
        val_name = ws.cell(r, 3).value
        if val_name and val_name != 'JUMLAH TOTAL':
            kg = ws.cell(r, 4).value or 0
            harga = ws.cell(r, 5).value or 0
            rp = kg * harga
            sec3.append({
                'no': ws.cell(r, 2).value,
                'nama': str(val_name).strip(),
                'kg': kg,
                'harga': harga,
                'rp': rp
            })

    # Sec IV (Stok)
    sec4 = []
    for r in range(28, 32):
        val_name = ws.cell(r, 7).value
        if val_name:
            stok_awal = ws.cell(r, 8).value or 0
            pembelian = ws.cell(r, 9).value or 0
            siap_jual = ws.cell(r, 10).value or 0
            penjualan = ws.cell(r, 11).value or 0
            susut = ws.cell(r, 12).value or 0
            stok_akhir = ws.cell(r, 13).value or 0
            # set default harga & rp
            harga_stok = 3900 if 'A18' in str(val_name) else (30000 if 'MAGNESIUM' in str(val_name) else (25000 if 'DCP' in str(val_name) else 4000))
            jumlah_rp = stok_akhir * harga_stok
            sec4.append({
                'no': ws.cell(r, 6).value,
                'nama': str(val_name).strip(),
                'stok_awal': stok_awal,
                'pembelian': pembelian,
                'siap_jual': siap_jual,
                'penjualan': penjualan,
                'susut': susut,
                'stok_akhir': stok_akhir,
                'harga': harga_stok,
                'jumlah_rp': jumlah_rp
            })

    db_all[name] = {
        'sec2': sec2,
        'sec3': sec3,
        'sec4': sec4
    }

with open('scratch/logistik_4sec_db.json', 'w', encoding='utf-8') as f:
    json.dump(db_all, f, indent=2)

print('Successfully generated scratch/logistik_4sec_db.json!')
