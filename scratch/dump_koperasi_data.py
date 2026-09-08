import openpyxl, json

def clean_num(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        return round(val, 2)
    try:
        s = str(val).replace('.', '').replace(',', '.').strip()
        return round(float(s), 2)
    except:
        return 0

wb_pen = openpyxl.load_workbook('LAPORAN BULANAN PENERIMAAN SUSU 2026.xlsx', data_only=True)
wb_peng = openpyxl.load_workbook('LAPORAN PENGGURUS 2026.xlsx', data_only=True)
wb_rek = openpyxl.load_workbook('LAPORAN PENERIMAAN DAN PENGELUARAN SUSU 2026.xlsx', data_only=True)

months = [
    ('JAN', 'Januari', 'Januari', 'Januari'),
    ('FEB', 'februari', 'februari', 'Februari'),
    ('MAR', 'maret', 'maret', 'Maret'),
    ('APRIL', 'April', 'April', 'April'),
    ('MEI', 'MEI', 'MEI', 'Mei'),
    ('JUNI', 'JUNI', 'JUNI', 'Juni'),
    ('JULI', 'JULI', 'JULI', 'juli.')
]

full_data = {}

for code, s_pen, s_peng, s_rek in months:
    m_data = {
        'penerimaan_anggota': [],
        'penerimaan_non_anggota': [],
        'pengeluaran_perusahaan': [],
        'pengeluaran_perorangan': [],
        'pengeluaran_lain_lain': [],
        'stok_rekap': {'stok_awal': 0, 'penerimaan': 0, 'persediaan': 0, 'pengeluaran': 0, 'stok_akhir': 0}
    }
    
    # 1. PENERIMAAN SUSU
    ws = wb_pen[s_pen]
    # Anggota rows (usually row 9 to 17 or so)
    for r in range(9, 35):
        kel = ws.cell(row=r, column=2).value
        kg = ws.cell(row=r, column=3).value
        rp = ws.cell(row=r, column=4).value
        
        wil = ws.cell(row=r, column=5).value
        w_kg = ws.cell(row=r, column=6).value
        w_rp = ws.cell(row=r, column=7).value
        
        if kel and str(kel).strip() and not str(kel).startswith('JUMLAH') and not str(kel).startswith('TOTAL'):
            m_data['penerimaan_anggota'].append({
                'kelompok': str(kel).strip(),
                'kg': clean_num(kg),
                'rp': clean_num(rp)
            })
            
        if wil and str(wil).strip() and not str(wil).startswith('JUMLAH') and not str(wil).startswith('TOTAL'):
            m_data['penerimaan_non_anggota'].append({
                'wilayah': str(wil).strip(),
                'kg': clean_num(w_kg),
                'rp': clean_num(w_rp)
            })

    # 2. PENGELUARAN SUSU
    ws_peng = wb_peng[s_peng]
    for r in range(9, 35):
        perus = ws_peng.cell(row=r, column=2).value
        peror = ws_peng.cell(row=r, column=3).value
        kg = ws_peng.cell(row=r, column=4).value
        hrg = ws_peng.cell(row=r, column=5).value
        rp = ws_peng.cell(row=r, column=6).value
        
        karena = ws_peng.cell(row=r, column=7).value
        l_kg = ws_peng.cell(row=r, column=8).value
        
        if perus and str(perus).strip() and str(perus).strip() in ['NESTLE', 'LOKAL', 'AGEN']:
            m_data['pengeluaran_perusahaan'].append({
                'nama': str(perus).strip(),
                'kg': clean_num(kg),
                'rp': clean_num(rp)
            })
        elif peror and str(peror).strip() and not str(peror).startswith('JUMLAH') and not str(peror).startswith('TOTAL'):
            m_data['pengeluaran_perorangan'].append({
                'nama': str(peror).strip(),
                'kg': clean_num(kg),
                'harga': clean_num(hrg),
                'rp': clean_num(rp)
            })
            
        if karena and str(karena).strip() and not str(karena).startswith('JUMLAH'):
            m_data['pengeluaran_lain_lain'].append({
                'karena': str(karena).strip(),
                'kg': clean_num(l_kg),
                'keterangan': str(karena).strip()
            })

    # 3. REKAP STOK
    if s_rek in wb_rek.sheetnames:
        ws_r = wb_rek[s_rek]
        for r in range(1, 40):
            v = str(ws_r.cell(row=r, column=1).value or '') + str(ws_r.cell(row=r, column=2).value or '')
            if 'Stoc' in v or 'Stok' in v or 'Stock' in v or 'Persediaan' in v or 'Pengeluaran' in v:
                val = ws_r.cell(row=r, column=3).value or ws_r.cell(row=r, column=4).value or ws_r.cell(row=r, column=5).value
                # print(f'{code} row {r}: {v} -> {val}')
                
    full_data[code] = m_data

print(json.dumps(full_data, indent=2))
