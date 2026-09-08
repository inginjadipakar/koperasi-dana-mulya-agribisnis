import openpyxl
import json

wb = openpyxl.load_workbook('LAPORAN BULANAN 2026 DEPOT SUSU.xlsx', data_only=True)

def to_float(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    try:
        s = str(val).replace('.', '').replace(',', '.').strip()
        return float(s)
    except:
        return 0.0

ignore_names = [
    'NAMA', 'KETERANGAN', 'LAPORAN', 'PEMBELIAN', 'PENJUALAN', 'TOTAL', 
    'JUMLAH SUSU SEGAR', 'JUMLAH', 'TOTAL PENJUALAN', 'TOTAL PEMBELIAN', 'TOTAL LAIN-LAIN', 'PROCESSING', 'ASAL'
]

months = ['JANUARI', 'FEBRUARI', 'MARET', 'APRIL', 'MEI', 'JUNI', 'JULI']
month_map = {
    'JANUARI': 'JAN',
    'FEBRUARI': 'FEB',
    'MARET': 'MAR',
    'APRIL': 'APRIL',
    'MEI': 'MEI',
    'JUNI': 'JUNI',
    'JULI': 'JULI'
}

depot_full_data = {}

for m in months:
    ws = wb[m if m in wb.sheetnames else m.lower()]
    m_code = month_map[m]
    
    pembelian = []
    penjualan = []
    lain_lain = []
    operasional = []
    
    # 1. Pembelian
    if m_code in ['JAN', 'FEB', 'MAR', 'APRIL', 'MEI']:
        c_kg = to_float(ws.cell(6, 3).value)
        c_liter = to_float(ws.cell(6, 5).value or (c_kg / 1.025))
        c_rp = to_float(ws.cell(6, 6).value or (c_kg * 9000))
        pembelian.append({
            'asal': 'PROCESSING',
            'kg': c_kg,
            'liter': round(c_liter, 2),
            'harga_per_kg': 9000,
            'total_rp': c_rp
        })
    else: # JUNI, JULI
        c_kg = to_float(ws.cell(7, 3).value)
        c_liter = to_float(ws.cell(7, 4).value)
        c_rp = to_float(ws.cell(7, 6).value)
        pembelian.append({
            'asal': 'PROCESSING',
            'kg': c_kg,
            'liter': round(c_liter, 2),
            'harga_per_kg': 9000,
            'total_rp': c_rp
        })

    # 2. Penjualan & Lain-lain & Operasional
    for r in range(1, 30):
        if m_code in ['JAN', 'FEB', 'MAR', 'APRIL', 'MEI']:
            c_nama = str(ws.cell(r, 7).value or '').strip()
            if c_nama and c_nama.upper() not in ignore_names:
                c_harga = to_float(ws.cell(r, 8).value)
                c_qty = to_float(ws.cell(r, 9).value)
                c_liter = to_float(ws.cell(r, 11).value if ws.cell(r, 11).value is not None else ws.cell(r, 10).value)
                c_rp = to_float(ws.cell(r, 12).value if ws.cell(r, 12).value is not None else ws.cell(r, 11).value)
                if (c_liter > 0 or c_rp > 0) and c_harga > 0:
                    penjualan.append({
                        'nama': c_nama.upper(),
                        'harga': c_harga,
                        'qty': c_qty,
                        'liter': round(c_liter, 2),
                        'rp': c_rp
                    })

            # Operasional
            c_ops_no = ws.cell(r, 18).value
            c_ops_nama = ws.cell(r, 19).value
            c_ops_rp = to_float(ws.cell(r, 22).value or ws.cell(r, 20).value)
            if isinstance(c_ops_no, (int, float)) and 1 <= c_ops_no <= 25 and c_ops_nama:
                operasional.append({
                    'no': int(c_ops_no),
                    'nama': str(c_ops_nama).strip(),
                    'rp': c_ops_rp
                })

            # Lain-lain
            c_lain = str(ws.cell(r, 13).value or '').strip()
            if c_lain and c_lain.upper() not in ignore_names:
                c_qty = to_float(ws.cell(r, 14).value)
                c_liter = to_float(ws.cell(r, 16).value)
                if c_qty > 0 or c_liter > 0:
                    lain_lain.append({
                        'keterangan': c_lain,
                        'qty': c_qty,
                        'liter': round(c_liter, 2)
                    })

        else: # JUNI, JULI
            c_nama = str(ws.cell(r, 7).value or '').strip()
            if c_nama and c_nama.upper() not in ignore_names:
                c_harga = to_float(ws.cell(r, 8).value)
                c_qty = to_float(ws.cell(r, 9).value)
                c_liter = to_float(ws.cell(r, 10).value)
                c_rp = to_float(ws.cell(r, 11).value)
                if (c_liter > 0 or c_rp > 0) and c_harga > 0:
                    penjualan.append({
                        'nama': c_nama.upper(),
                        'harga': c_harga,
                        'qty': c_qty,
                        'liter': round(c_liter, 2),
                        'rp': c_rp
                    })

            # Operasional
            c_ops_no = ws.cell(r, 17).value
            c_ops_nama = ws.cell(r, 18).value
            c_ops_rp = to_float(ws.cell(r, 21).value or ws.cell(r, 19).value)
            if isinstance(c_ops_no, (int, float)) and 1 <= c_ops_no <= 25 and c_ops_nama:
                operasional.append({
                    'no': int(c_ops_no),
                    'nama': str(c_ops_nama).strip(),
                    'rp': c_ops_rp
                })

            # Lain-lain
            c_lain = str(ws.cell(r, 12).value or '').strip()
            if c_lain and c_lain.upper() not in ignore_names:
                c_qty = to_float(ws.cell(r, 13).value)
                c_liter = to_float(ws.cell(r, 14).value)
                if c_qty > 0 or c_liter > 0:
                    lain_lain.append({
                        'keterangan': c_lain,
                        'qty': c_qty,
                        'liter': round(c_liter, 2)
                    })

    # Stok rekap
    if m_code in ['JAN', 'FEB', 'MAR', 'APRIL', 'MEI']:
        rekap_stok = {
            'stok_awal': to_float(ws.cell(26, 8).value),
            'penerimaan': to_float(ws.cell(27, 7).value),
            'persediaan': to_float(ws.cell(29, 8).value),
            'pengeluaran': to_float(ws.cell(30, 8).value),
            'stok_akhir': to_float(ws.cell(32, 8).value or ws.cell(31, 8).value),
            'nb': str(ws.cell(34, 1).value or '').strip()
        }
    else:
        stok_awal_val = to_float(ws.cell(33, 7).value)
        penerimaan_val = to_float(ws.cell(34, 7).value)
        pengeluaran_val = to_float(ws.cell(28, 10).value)
        rekap_stok = {
            'stok_awal': stok_awal_val,
            'penerimaan': penerimaan_val,
            'persediaan': round(stok_awal_val + penerimaan_val, 2),
            'pengeluaran': pengeluaran_val,
            'stok_akhir': round(stok_awal_val + penerimaan_val - pengeluaran_val, 2),
            'nb': f"Sisa botol: {ws.cell(33, 4).value or 0}, sisa liter: {ws.cell(33, 6).value or 0}"
        }

    depot_full_data[m_code] = {
        'pembelian': pembelian,
        'penjualan': penjualan,
        'lain_lain': lain_lain,
        'operasional': operasional,
        'stok': rekap_stok
    }

with open('scratch/depot_full_final.json', 'w') as f:
    json.dump(depot_full_data, f, indent=2)

print("--- REVISED DEPOT SUMMARY ---")
for m, obj in depot_full_data.items():
    tot_sal_ltr = sum(p['liter'] for p in obj['penjualan'])
    tot_sal_rp = sum(p['rp'] for p in obj['penjualan'])
    print(f"{m:6s} | Penj Rows: {len(obj['penjualan']):2d} | Total Ltr: {tot_sal_ltr:10,.1f} | Total Rp: Rp {tot_sal_rp:14,.0f}")
