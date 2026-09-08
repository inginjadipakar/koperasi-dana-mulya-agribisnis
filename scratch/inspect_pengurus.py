#!/usr/bin/env python3
from pathlib import Path
import openpyxl

PROJECT_ROOT = Path(r"c:\Users\HP\Desktop\KKN SIE ACARA\Proker UMKM\Koperasidanamulya")
filepath = PROJECT_ROOT / "LAPORAN PENGGURUS 2026.xlsx"

wb = openpyxl.load_workbook(filepath, data_only=True)

print("============================================================")
print("INSPEKSI LAPORAN PENGGURUS 2026.xlsx (EXCEL ASLI PUSAT)")
print("============================================================\n")

print("Daftar Sheet:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- SHEET: {sheetname} ({ws.max_row} baris, {ws.max_column} kolom) ---")
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cleaned = [str(c) if c is not None else "" for c in row]
        if any(c != "" for c in cleaned):
            print(f"  Baris {r_idx:2d}: {cleaned[:8]}")
