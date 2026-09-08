#!/usr/bin/env python3
import json
from pathlib import Path
import openpyxl

PROJECT_ROOT = Path(r"c:\Users\HP\Desktop\KKN SIE ACARA\Proker UMKM\Koperasidanamulya")
OUTPUT_JS = PROJECT_ROOT / "js" / "excel_data.js"

EXCEL_FILES = [
    "LAPORAN BULANAN 2026 DEPOT SUSU.xlsx",
    "LAPORAN BULANAN PENERIMAAN SUSU 2026.xlsx",
    "LAPORAN PENERIMAAN DAN PENGELUARAN SUSU 2026.xlsx",
    "LAPORAN PENGGURUS 2026.xlsx"
]

all_data = {}

for filename in EXCEL_FILES:
    filepath = PROJECT_ROOT / filename
    if not filepath.exists():
        print("Not found:", filename)
        continue
    
    print("Processing:", filename)
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_data[filename] = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_rows = []
        for row in ws.iter_rows(values_only=True):
            # Clean row values
            cleaned_row = []
            for cell in row:
                if cell is None:
                    cleaned_row.append("")
                else:
                    cleaned_row.append(str(cell))
            # Keep non-empty rows
            if any(c != "" for c in cleaned_row):
                sheet_rows.append(cleaned_row)
        
        all_data[filename][sheet_name] = sheet_rows

js_content = f"const ORIGINAL_EXCEL_DATA = {json.dumps(all_data, indent=2, ensure_ascii=False)};"
OUTPUT_JS.write_text(js_content, encoding="utf-8")
print("Saved to:", OUTPUT_JS)
